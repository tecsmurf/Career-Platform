"""
Isolated resume parser — runs as a separate, resource-limited process.

PDF and DOCX files are attacker-controllable: a few hundred KB can decompress
into gigabytes (zip / flate bombs). Parsing therefore happens in a child
process that caps its own memory and CPU time and that the server kills if it
runs past its deadline, so a hostile file can only ever kill this worker.

Protocol: file bytes on stdin, kind ("pdf" | "docx") in argv[1];
one JSON object on stdout: {"ok": true, "text": "..."} or {"ok": false, "error": "..."}.
This module deliberately imports nothing from the app (no settings, no DB).
"""
import io
import json
import sys

MAX_PAGES = 15
MAX_CHARS = 40_000
MEMORY_LIMIT = 512 * 1024 * 1024
CPU_SECONDS = 20


def _limit_resources() -> None:
    try:
        import resource
    except ImportError:  # Windows: no rlimits; the parent's timeout still applies
        return
    for limit, value in ((resource.RLIMIT_AS, MEMORY_LIMIT), (resource.RLIMIT_CPU, CPU_SECONDS)):
        try:
            resource.setrlimit(limit, (value, value))
        except (ValueError, OSError):
            pass


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ValueError("encrypted")
    out = []
    total = 0
    for page in reader.pages[:MAX_PAGES]:
        text = page.extract_text() or ""
        out.append(text)
        total += len(text)
        if total > MAX_CHARS:
            break
    return "\n".join(out)


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = []
    total = 0
    for p in document.paragraphs:
        lines.append(p.text)
        total += len(p.text) + 1
        if total > MAX_CHARS:
            break
    for table in document.tables:
        for row in table.rows:
            lines.append("  ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def main() -> None:
    _limit_resources()
    kind = sys.argv[1] if len(sys.argv) > 1 else ""
    data = sys.stdin.buffer.read()
    try:
        if kind == "pdf":
            text = _pdf(data)
        elif kind == "docx":
            text = _docx(data)
        else:
            raise ValueError("kind")
        result = {"ok": True, "text": text[: MAX_CHARS * 2]}
    except MemoryError:
        result = {"ok": False, "error": "too_large"}
    except ValueError as exc:
        result = {"ok": False, "error": "encrypted" if str(exc) == "encrypted" else "unreadable"}
    except Exception:  # any parser failure → "unreadable" (never echo parser internals)
        result = {"ok": False, "error": "unreadable"}
    sys.stdout.write(json.dumps(result))


if __name__ == "__main__":
    main()
